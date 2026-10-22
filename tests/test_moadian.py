"""آزمون‌های فاز ۵ — سامانه‌ی مودیان: ساخت صورتحساب الکترونیکی، صف ارسال و گزارش ارزش افزوده."""
from __future__ import annotations

import json
from io import BytesIO

from openpyxl import load_workbook

from app import dates as jd

GOOD_STUFF_ID = "2717030010101"      # شناسه‌ی نمونه (فرمت درست: ۱۳ رقم)


def _settings(client, accountant, active=True, **over):
    body = {"taxid": "A" + "1" * 21, "economic_code": "14009998888", "seller_name": "آژانس آسمان‌گشت",
            "person_type": 2, "send_method": "motamed", "motamed_name": "شرکت معتمد نمونه",
            "environment": "sandbox", "is_active": active, "postal_code": "1234567890",
            "phone": "02188880000", "address": "تهران"}
    body.update(over)
    r = client.post("/api/moadian/settings", headers=accountant, json=body)
    assert r.status_code == 200, r.text
    return body


def stuff_id_for(index: int) -> str:
    """شناسه‌ی ۱۳ رقمی یکتا برای هر نوع خدمت (شناسه‌ی اختصاصی در هر خدمت یک‌بار)."""
    return f"2717030010{index:03d}"


def _services_ready(client, accountant):
    """همه‌ی شناسه‌های کالا/خدمت را وارد و تأیید می‌کند."""
    rows = client.get("/api/moadian/service-map", headers=accountant).json()["rows"]
    for r in rows:
        resp = client.patch(f"/api/moadian/service-map/{r['id']}", headers=accountant,
                            json={"stuff_id": stuff_id_for(r["id"]), "is_confirmed": True})
        assert resp.status_code == 200, resp.text
    return rows


def _party(client, accountant, name, **over):
    body = {"name": name, "kind": "customer"}
    body.update(over)
    r = client.post("/api/parties", headers=accountant, json=body)
    assert r.status_code == 200, r.text
    party = r.json()
    if any(k in over for k in ("national_id", "economic_code", "person_type")):
        # اطلاعات مالیاتی بعد از ساخت شخص هم قابل ثبت است
        r2 = client.patch(f"/api/parties/{party['id']}", headers=accountant, json=over)
        if r2.status_code == 200:
            party = r2.json()
    return party


def _invoice(client, accountant, party_id, lines, issue=True, date=None):
    r = client.post("/api/invoices", headers=accountant, json={
        "party_id": party_id, "date": date or jd.jalali_str(jd.jtoday()), "issue": issue,
        "description": "آزمون مودیان", "lines": lines})
    assert r.status_code == 200, r.text
    return r.json()


def _build(client, accountant, owner="accountant"):
    r = client.post("/api/moadian/build", headers=accountant, json={
        "frm": jd.jalali_str(jd.jtoday()), "to": jd.jalali_str(jd.jtoday()),
        "include_cancelled": True})
    assert r.status_code == 200, r.text
    return r.json()


# ------------------------------------------------------------ صفحه و دسترسی
def test_moadian_page_opens_and_shows_demo_queue(client, accountant):
    """صفحه‌ی مودیان باز می‌شود و داده‌ی نمونه‌ی آژانس را نشان می‌دهد."""
    ov = client.get("/api/moadian/overview", headers=accountant)
    assert ov.status_code == 200, ov.text
    data = ov.json()
    assert data["settings"]["is_active"] is False          # در نمونه، ارسال واقعی فعال نیست
    assert all(k in data["summary"]["counts"] for k in ("draft", "ready", "accepted"))
    q = client.get("/api/moadian/queue", headers=accountant)
    assert q.status_code == 200
    assert q.json()["rows"], "صف نمونه باید صورتحساب داشته باشد"
    assert any(r["status"] == "accepted" for r in q.json()["rows"])


def test_employee_cannot_see_tax_system(client, employee):
    assert client.get("/api/moadian/overview", headers=employee).status_code == 403


# ------------------------------------------------------------ تنظیمات
def test_settings_reject_wrong_taxid_length(client, accountant):
    r = client.post("/api/moadian/settings", headers=accountant, json={
        "taxid": "12345", "economic_code": "14009998888", "send_method": "motamed",
        "motamed_name": "شرکت معتمد نمونه"})
    assert r.status_code == 400
    assert "۲۲" in r.json()["detail"]


def test_settings_with_motamed_need_motamed_name(client, accountant):
    r = client.post("/api/moadian/settings", headers=accountant, json={
        "taxid": "A" + "1" * 21, "economic_code": "14009998888", "send_method": "motamed",
        "motamed_name": ""})
    assert r.status_code == 400
    assert "معتمد" in r.json()["detail"]


def test_invoice_without_stuff_id_is_not_buildable(client, accountant):
    """آزمون نمونه‌ی فاز ۵: تا شناسه‌ی کالا/خدمت نباشد، صورتحساب «آماده‌ی ارسال» نمی‌شود."""
    _settings(client, accountant)
    # شناسه‌ی خدمت «اقامت» را موقتاً برمی‌داریم تا فاکتور حتماً ناقص بماند
    rows = client.get("/api/moadian/service-map", headers=accountant).json()["rows"]
    hotel = next(r for r in rows if r["service_type"] == "hotel")
    assert client.patch(f"/api/moadian/service-map/{hotel['id']}", headers=accountant,
                        json={"stuff_id": "", "is_confirmed": False}).status_code == 200
    party = _party(client, accountant, "مشتری بدون شناسه")
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "hotel", "unit_price": 100_000_000, "qty": 1, "taxable": True}])
    res = _build(client, accountant)
    assert res["created"] >= 1
    row = next(r for r in client.get("/api/moadian/queue", headers=accountant).json()["rows"]
               if r["invoice_number"] == inv["number"])
    assert row["status"] == "draft"
    assert any("شناسه‌ی کالا/خدمت" in p for p in row["problems"]), row["problems"]
    # جای شناسه برمی‌گردد تا بقیه‌ی آزمون‌ها سالم بمانند
    assert client.patch(f"/api/moadian/service-map/{hotel['id']}", headers=accountant,
                        json={"stuff_id": stuff_id_for(hotel["id"]),
                              "is_confirmed": True}).status_code == 200


def test_after_confirming_stuff_ids_invoice_becomes_ready(client, accountant):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مشتری آماده")
    _invoice(client, accountant, party["id"], [
        {"service_type": "tour", "unit_price": 100_000_000, "qty": 1, "taxable": True}])
    _build(client, accountant)
    rows = client.get("/api/moadian/queue?status=draft", headers=accountant).json()["rows"]
    for r in rows:
        assert client.post(f"/api/moadian/queue/{r['id']}/rebuild",
                           headers=accountant).status_code == 200
    ready = client.get("/api/moadian/queue?status=ready", headers=accountant).json()["rows"]
    assert ready, "بعد از وارد کردن شناسه‌ها باید حداقل یک صورتحساب آماده شود"


# ------------------------------------------------------------ ساختار صورتحساب
def test_payload_matches_official_structure_and_totals(client, accountant):
    """ساختار صورتحساب همان فیلدهای رسمی سامانه‌ی مودیان است و جمع‌ها درست‌اند."""
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "شرکت خریدار آزمایشی", national_id="0070000001",
                   economic_code="14008887777", person_type=2, postal_code="1998765432")
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "tour", "unit_price": 100_000_000, "qty": 1, "taxable": True},
        {"service_type": "visa", "unit_price": 50_000_000, "qty": 2, "discount": 10_000_000,
         "taxable": True}])
    _build(client, accountant)
    row = next(r for r in client.get("/api/moadian/queue", headers=accountant).json()["rows"]
               if r["invoice_number"] == inv["number"])
    assert row["status"] == "ready", row["problems"]
    d = client.get(f"/api/moadian/queue/{row['id']}", headers=accountant).json()
    h = d["payload"]["header"]
    for field in ("taxid", "indatim", "inty", "inno", "inp", "ins", "tins", "tob", "bid", "tinb",
                  "sbc", "bpc", "tprdis", "tdis", "tadis", "tvam", "todam", "tbill", "setm",
                  "cap", "insp", "tvop"):
        assert field in h, f"فیلد {field} در سرآمد نیست"
    assert h["inty"] == 1                      # خریدار حقوقی با شناسه → نوع اول
    assert h["ins"] == 1                       # صورتحساب اصلی
    assert h["inp"] == 1                       # الگوی «فروش»
    assert h["tob"] == 2 and h["tinb"] == "14008887777" and h["bid"] == "0070000001"
    assert h["tprdis"] == 200_000_000          # ۱۰۰م + ۲×۵۰م
    assert h["tdis"] == 10_000_000
    assert h["tadis"] == 190_000_000
    # مالیات با نرخ ۱۰٪ مرداد ۱۴۰۵
    assert h["tvam"] == 19_000_000, h["tvam"]
    assert h["tbill"] == h["tadis"] + h["tvam"]
    assert h["setm"] == 2 and h["cap"] == 0 and h["insp"] == h["tbill"]   # نسیه
    body = d["payload"]["body"]
    assert len(body) == 2
    assert all(len(l["sstid"]) == 13 for l in body)
    assert sum(l["prdis"] for l in body) == h["tprdis"]
    assert sum(l["adis"] for l in body) == h["tadis"]
    assert sum(l["vam"] for l in body) == h["tvam"]
    assert all(l["vra"] == 0.1 for l in body)  # نرخ از جدول نرخ‌ها، نه ثابت در کد


def test_cash_customer_without_ids_becomes_type_two(client, accountant):
    """آزمون نمونه: مشتری بدون کد ملی → نوع دوم (بدون اعتبار مالیاتی) و بدون ایراد."""
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مسافر نقدی بی‌شناسه")
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "hotel", "unit_price": 30_000_000, "qty": 1, "taxable": True}])
    _build(client, accountant)
    row = next(r for r in client.get("/api/moadian/queue", headers=accountant).json()["rows"]
               if r["invoice_number"] == inv["number"])
    assert row["type"] == 2, row["problems"]


def test_paid_invoice_gets_cash_settlement(client, accountant):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مشتری نقدی پرداخت‌کننده")
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "visa", "unit_price": 20_000_000, "qty": 1, "taxable": True}])
    bank = next(a for a in client.get("/api/accounts", headers=accountant).json()["rows"]
                if a["code"] == "110201")
    pay = client.post("/api/settlements", headers=accountant, json={
        "direction": "in", "party_id": party["id"], "account_id": bank["id"],
        "amount": inv["total"], "date": jd.jalali_str(jd.jtoday()),
        "description": f"تسویه‌ی فاکتور {inv['number']}", "issue": True})
    assert pay.status_code == 200, pay.text
    _build(client, accountant)
    row = next(r for r in client.get("/api/moadian/queue", headers=accountant).json()["rows"]
               if r["invoice_number"] == inv["number"])
    d = client.get(f"/api/moadian/queue/{row['id']}", headers=accountant).json()
    h = d["payload"]["header"]
    assert h["setm"] == 1 and h["cap"] == h["tbill"] and h["insp"] == 0


# ------------------------------------------------------------ بلیت هواپیما
def test_ticket_uses_flight_pattern_and_is_type_one(client, accountant):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مسافر پرواز", national_id="0069998877",
                   person_type=1)
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "ticket", "unit_price": 70_000_000, "qty": 1, "taxable": False,
         "flight_type": 2, "description": "بلیت تهران — استانبول"}])
    _build(client, accountant)
    row = next(r for r in client.get("/api/moadian/queue", headers=accountant).json()["rows"]
               if r["invoice_number"] == inv["number"])
    assert row["pattern"] == 6, row["problems"]      # الگوی ۶ = بلیت هواپیما
    assert row["type"] == 1, row["problems"]         # بلیت همیشه نوع اول
    d = client.get(f"/api/moadian/queue/{row['id']}", headers=accountant).json()
    assert d["payload"]["header"]["ft"] == 2         # پرواز خارجی
    assert d["payload"]["body"][0]["flight_type"] == 2
    assert d["payload"]["header"]["tvam"] == 0       # بلیت معاف از مالیات


def test_flight_type_can_be_registered_later_but_not_after_sending(client, accountant):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مسافر بی‌نوع‌پرواز", national_id="0069998876")
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "ticket", "unit_price": 70_000_000, "qty": 1, "taxable": False,
         "description": "بلیت بدون نوع پرواز"}])
    _build(client, accountant)
    rows = client.get("/api/moadian/queue?status=draft", headers=accountant).json()["rows"]
    mine = [r for r in rows if r["invoice_number"] == inv["number"]]
    assert mine and any("نوع پرواز" in p for p in mine[0]["problems"])
    detail_before = client.get(f"/api/moadian/queue/{mine[0]['id']}", headers=accountant).json()
    assert detail_before["ticket_lines"], "ردیف بلیت باید در جزئیات صورتحساب باشد"
    line_id = detail_before["ticket_lines"][0]["line_id"]
    fix = client.post(f"/api/invoices/{inv['id']}/flight-type", headers=accountant,
                      json={"line_id": line_id, "flight_type": 1})
    assert fix.status_code == 200, fix.text
    client.post(f"/api/moadian/queue/{mine[0]['id']}/rebuild", headers=accountant)
    after = client.get(f"/api/moadian/queue/{mine[0]['id']}", headers=accountant).json()
    assert not after["problems"], after["problems"]
    assert after["payload"]["header"]["ft"] == 1

    # بعد از ارسال، دیگر تغییر نوع پرواز معنی ندارد
    client.post(f"/api/moadian/queue/{mine[0]['id']}/result", headers=accountant,
                json={"status": "sent", "reference_number": "REF-1"})
    again = client.post(f"/api/invoices/{inv['id']}/flight-type", headers=accountant,
                        json={"line_id": line_id, "flight_type": 2})
    assert again.status_code == 400
    assert "ارسال" in again.json()["detail"]


# ------------------------------------------------------------ صف و نتیجه‌ی ارسال
def test_build_is_idempotent(client, accountant):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مشتری تکرار")
    _invoice(client, accountant, party["id"], [
        {"service_type": "tour", "unit_price": 10_000_000, "qty": 1, "taxable": True}])
    first = _build(client, accountant)
    second = _build(client, accountant)
    assert second["created"] == 0 and second["skipped"] >= 1
    assert first["created"] >= 1


def test_accepted_needs_taxid_and_status_stays(client, accountant):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مشتری پذیرش")
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "tour", "unit_price": 10_000_000, "qty": 1, "taxable": True}])
    _build(client, accountant)
    row = next(r for r in client.get("/api/moadian/queue", headers=accountant).json()["rows"]
               if r["invoice_number"] == inv["number"])
    no_taxid = client.post(f"/api/moadian/queue/{row['id']}/result", headers=accountant,
                           json={"status": "accepted", "reference_number": "REF-9"})
    assert no_taxid.status_code == 400
    ok = client.post(f"/api/moadian/queue/{row['id']}/result", headers=accountant,
                     json={"status": "accepted", "reference_number": "REF-9",
                           "taxid": "B" + "2" * 21})
    assert ok.status_code == 200, ok.text
    detail = client.get(f"/api/moadian/queue/{row['id']}", headers=accountant).json()
    assert detail["status"] == "accepted" and detail["taxid"] == "B" + "2" * 21
    # عقب بردن صورتحساب پذیرفته‌شده ممنوع است
    back = client.post(f"/api/moadian/queue/{row['id']}/result", headers=accountant,
                       json={"status": "sent", "reference_number": "REF-9"})
    assert back.status_code == 400


def test_batch_file_and_single_file_are_valid_json(client, accountant):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مشتری فایل")
    _invoice(client, accountant, party["id"], [
        {"service_type": "tour", "unit_price": 12_000_000, "qty": 1, "taxable": True}])
    _build(client, accountant)
    batch = client.get("/api/moadian/batch-file?status=ready", headers=accountant)
    assert batch.status_code == 200
    data = json.loads(batch.content.decode("utf-8"))
    assert data["بسته‌ها"] and "header" in data["بسته‌ها"][0]["packet"]
    row = client.get("/api/moadian/queue", headers=accountant).json()["rows"][0]
    single = client.get(f"/api/moadian/queue/{row['id']}/file", headers=accountant)
    assert single.status_code == 200
    assert json.loads(single.content.decode("utf-8"))["header"]["inty"] in (1, 2)


def test_draft_invoice_has_no_file(client, accountant):
    """صورتحساب ناقص نباید فایل ارسالی داشته باشد."""
    _settings(client, accountant)
    # شناسه‌ی خدمت «بیمه» را موقتاً برمی‌داریم تا فاکتور ناقص بماند
    rows = client.get("/api/moadian/service-map", headers=accountant).json()["rows"]
    ins = next(r for r in rows if r["service_type"] == "insurance")
    assert client.patch(f"/api/moadian/service-map/{ins['id']}", headers=accountant,
                        json={"stuff_id": ""}).status_code == 200
    party = _party(client, accountant, "مشتری بدون شناسه‌ی خدمت")
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "insurance", "unit_price": 15_000_000, "qty": 1, "taxable": True}])
    _build(client, accountant)
    row = next(r for r in client.get("/api/moadian/queue", headers=accountant).json()["rows"]
               if r["invoice_number"] == inv["number"])
    assert row["status"] == "draft"
    assert client.get(f"/api/moadian/queue/{row['id']}/file",
                      headers=accountant).status_code == 400
    # جای خودش برمی‌گردد تا بقیه‌ی آزمون‌ها سالم بمانند
    assert client.patch(f"/api/moadian/service-map/{ins['id']}", headers=accountant,
                        json={"stuff_id": stuff_id_for(ins["id"]),
                              "is_confirmed": True}).status_code == 200


def test_cancelled_invoice_creates_subject_three(client, accountant):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مشتری ابطال", national_id="0069998875")
    inv = _invoice(client, accountant, party["id"], [
        {"service_type": "tour", "unit_price": 25_000_000, "qty": 1, "taxable": True}])
    _build(client, accountant)
    row = next(r for r in client.get("/api/moadian/queue", headers=accountant).json()["rows"]
               if r["invoice_number"] == inv["number"])
    client.post(f"/api/moadian/queue/{row['id']}/result", headers=accountant,
                json={"status": "accepted", "reference_number": "REF-7", "taxid": "C" + "3" * 21})
    cancel = client.post(f"/api/invoices/{inv['id']}/cancel", headers=accountant,
                         json={"reason": "آزمون ابطال"})
    assert cancel.status_code == 200, cancel.text
    _build(client, accountant)
    rows = client.get("/api/moadian/queue", headers=accountant).json()["rows"]
    mine = [r for r in rows if r["invoice_number"] == inv["number"] and r["subject"] == 3]
    assert mine, "فاکتور ابطالی باید صورتحساب ابطالی بسازد"
    detail = client.get(f"/api/moadian/queue/{mine[0]['id']}", headers=accountant).json()
    assert detail["payload"]["header"]["ins"] == 3
    assert detail["payload"]["header"]["irtaxid"] == "C" + "3" * 21


# ------------------------------------------------------------ گزارش ارزش افزوده
def test_vat_report_separates_exempt_tickets(client, accountant):
    """آزمون نمونه: گزارش فصلی، فروش مشمول را از معاف جدا می‌کند."""
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مشتری گزارش مالیات")
    # today = jd.jalali_str(jd.jtoday())
    _invoice(client, accountant, party["id"], [
        {"service_type": "tour", "unit_price": 100_000_000, "qty": 1, "taxable": True}])
    _invoice(client, accountant, party["id"], [
        {"service_type": "ticket", "unit_price": 40_000_000, "qty": 1, "taxable": False,
         "flight_type": 1}])
    jy, jm, _ = jd.jalali_parts(jd.jtoday())
    rep = client.get(f"/api/reports/vat?jy={jy}&jm={jm}", headers=accountant)
    assert rep.status_code == 200, rep.text
    d = rep.json()
    assert d["taxable_sales"] >= 100_000_000
    assert d["exempt_sales"] >= 40_000_000
    assert d["output_vat"] >= 10_000_000
    assert d["payable_vat"] >= 0
    assert d["quarter"] == (jm - 1) // 3 + 1
    assert d["rows"], "فهرست فاکتورهای فصل باید خالی نباشد"


def test_vat_report_of_previous_quarter_is_separate(client, accountant):
    jy, jm, _ = jd.jalali_parts(jd.jtoday())
    prev = 1 if jm > 3 else 4                      # یک فصل دیگر
    rep = client.get(f"/api/reports/vat?jy={jy if jm > 3 else jy - 1}&jm={prev}",
                     headers=accountant).json()
    today_rep = client.get(f"/api/reports/vat?jy={jy}&jm={jm}", headers=accountant).json()
    assert rep["rows"] != today_rep["rows"]


def test_vat_and_queue_excel_exports_work(client, accountant, admin):
    _settings(client, accountant)
    _services_ready(client, accountant)
    party = _party(client, accountant, "مشتری خروجی اکسل")
    _invoice(client, accountant, party["id"], [
        {"service_type": "tour", "unit_price": 10_000_000, "qty": 1, "taxable": True}])
    _build(client, accountant)
    qx = client.get("/api/export/moadian-queue.xlsx", headers=accountant)
    vx = client.get("/api/export/vat.xlsx", headers=accountant)
    for r in (qx, vx):
        assert r.status_code == 200, r.text
        assert r.content[:2] == b"PK", "فایل اکسل سالم نیست"
    wb = load_workbook(BytesIO(qx.content), read_only=True, data_only=True)
    assert wb.sheetnames == ["سرآمد", "اقلام", "پرداخت", "کنترل", "راهنما"]
    assert "taxid" in [cell.value for cell in wb["سرآمد"][1]]
    guide = " ".join(str(cell.value or "") for row in wb["راهنما"].iter_rows() for cell in row)
    assert "نه قالب رسمی بارگذاری" in guide
    assert "JSON" in guide
    wb.close()


def test_service_map_rejects_duplicate_and_short_ids(client, accountant):
    rows = client.get("/api/moadian/service-map", headers=accountant).json()["rows"]
    first, second, third = rows[0]["id"], rows[1]["id"], rows[2]["id"]
    short = client.patch(f"/api/moadian/service-map/{first}", headers=accountant,
                         json={"stuff_id": "12345"})
    assert short.status_code == 400 and "۱۳ رقم" in short.json()["detail"]
    used = stuff_id_for(first)
    ok = client.patch(f"/api/moadian/service-map/{first}", headers=accountant,
                      json={"stuff_id": used})
    assert ok.status_code == 200
    dup = client.patch(f"/api/moadian/service-map/{second}", headers=accountant,
                       json={"stuff_id": used})
    assert dup.status_code == 400 and "قبلاً" in dup.json()["detail"]
    # تأیید بدون شناسه ممنوع است
    cleared = client.patch(f"/api/moadian/service-map/{third}", headers=accountant,
                           json={"stuff_id": ""})
    assert cleared.status_code == 200
    empty = client.patch(f"/api/moadian/service-map/{third}", headers=accountant,
                         json={"is_confirmed": True})
    assert empty.status_code == 400
